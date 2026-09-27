#!/usr/bin/env python3
"""Yuj work item: count cells for one single-cell accession."""

import functools
import json
import logging
import os
import re
import shutil
import signal
import socket
import sys
import tarfile
from pathlib import Path
from typing import Any

from helpers import (
    FileTooLargeError,
    download_file,
    download_file_ftp,
    download_ftp_head,
    download_http_head,
    ftp_data_stream,
)
from sc_enrich.deep._process import (
    FILE_TYPE_PRIORITY,
    HEADER_STREAMABLE,
    _check_disk_space,
    _get_sc_count,
    _member_head,
    _resolve_file_type,
    count_h5_dims_via_http,
    count_tar_members_via_stream,
    count_tar_via_stream,
    gene_count_for,
    h5ad_obs_columns_local,
    h5ad_obs_columns_via_http,
    is_cell_major,
    last_mtx_nnz,
    last_rds_class,
    process_tar_with_genes,
    rds_columns,
)
from sc_enrich.deep.counter_common import DownloadLink, FileType, RdsReadError

logger = logging.getLogger(__name__)

_FAILURES_LOG = "rds_read_failures.jsonl"


def _record_rds_failure(
    output_dir: Path, accession: str, url: str, reason: str
) -> None:
    # one worker per host, so the append is race-free; per-host logs merge on yuj pull
    rec = {
        "host": socket.gethostname(),
        "accession": accession,
        "url": url,
        "reason": reason,
    }
    with (output_dir / _FAILURES_LOG).open("a", encoding="utf-8") as f:
        f.write(json.dumps(rec) + "\n")


_CHUNK_SIZE = 10 * 1024 * 1024
_RETRIES = 5
_TIMEOUT = 30
_MAX_TIMEOUT = 300

_MAX_FILE_BYTES = 30 * 1024**3

_MTX_HEAD_BYTES = 1 * 1024 * 1024

_META_HEAD_BYTES = 256 * 1024
_MAX_META_COLS = 40
_META_EXCLUDE = (
    ".mtx",
    "barcodes.tsv",
    "features.tsv",
    "genes.tsv",
    ".h5",
    ".hdf5",
)
_META_EXTS = (".csv", ".tsv", ".txt", ".xlsx", ".xls")
_META_TABLE_EXTS = (".xlsx", ".xls")
_EMBEDDED_METADATA_FORMATS = ("rds", "h5ad")

# Raw/unfiltered 10x matrices report whole barcode whitelists as columns;
# detect exact whitelist sizes or low nnz/cells.
_WHITELIST_SIZES = frozenset(
    {737280, 6794880, 4526976, 3571200, 6390720, 1000000, 900000},
)
_MIN_NNZ_PER_CELL = 20
# For low-feature matrices, compare density with the feature count. The 0.5
# cutoff separates observed filtered (>=0.8) and raw (<0.2) matrices.
_MIN_LOW_FEATURE_DENSITY = 0.5
_MAX_TARS_SCANNED = 2
_MAX_TAR_META_MEMBERS = 50
_MAX_TAR_MEMBERS_WALKED = 4000
_BARCODE_COL_HINTS = (
    "barcode",
    "cell_id",
    "cellid",
    "cell.id",
    "cell_name",
    "cellname",
    "cell_barcode",
)
# Count named biological annotations only; numeric clusters inflate has_celltype.
_CELLTYPE_COL_HINTS = (
    "celltype",
    "cell_type",
    "cell.type",
    "cell_class",
    "annotation",
    "annot",
    "subtype",
    "cell_subtype",
    "cell_label",
    "majortype",
    "major_type",
    "minortype",
    "minor_type",
    "lineage",
    "population",
    "cell_ontology",
    "cellontology",
)

_DEMOGRAPHIC_CATEGORIES = ("sex", "age", "race_ancestry", "bmi", "smoking")
_META_FIELD_HINTS: dict[str, tuple[str, ...]] = {
    "subject": (
        "subject",
        "individual",
        "patient",
        "participant",
        "donor",
        "orig.ident",
        "mouse",
        "animal",
    ),
    "sample_id": (
        "sample",
        "sample_id",
        "sampleid",
        "biosample",
        "specimen",
        "aliquot",
        "gsm",
    ),
    "celltype": _CELLTYPE_COL_HINTS,
    "cluster": (
        "cluster",
        "seurat_clusters",
        "leiden",
        "louvain",
        "ident",
        "subcluster",
        "resolution",
    ),
    "sex": ("sex", "gender"),
    "age": ("age",),
    "race_ancestry": ("race", "ancestry", "ethnicity", "ethnic", "population"),
    "bmi": ("bmi", "body_mass_index", "obese", "obesity"),
    "smoking": ("smoking", "smoker", "smoke", "tobacco", "pack_years", "packyears"),
    "tissue": (
        "tissue",
        "organ",
        "region",
        "anatomical",
        "site",
        "compartment",
        "source_name",
        "biopsy",
        "isolation_source",
        "organism_part",
    ),
    "disease": ("disease", "diagnosis", "phenotype", "pathology", "lesion"),
    "disease_stage": (
        "stage",
        "grade",
        "severity",
        "tnm",
        "metastatic",
        "metastasis",
        "remission",
        "relapse",
        "progression",
    ),
    "condition": (
        "condition",
        "group",
        "status",
        "case",
        "control",
        "healthy",
        "normal",
        "disease_status",
        "case_control",
    ),
    "treatment": (
        "treatment",
        "treated",
        "untreated",
        "drug",
        "therapy",
        "intervention",
        "exposure",
        "stim",
        "stimulation",
        "infection",
        "infected",
        "dose",
        "compound",
        "vehicle",
        "perturbation",
    ),
    "timepoint": (
        "timepoint",
        "time_point",
        "time",
        "day",
        "week",
        "month",
        "hour",
        "baseline",
        "followup",
        "follow_up",
        "visit",
        "longitudinal",
    ),
    "genotype": (
        "genotype",
        "mutation",
        "variant",
        "allele",
        "strain",
        "knockout",
        "wildtype",
        "transgenic",
        "ko",
        "wt",
        "tg",
        "carrier",
    ),
    "medication": (
        "medication",
        "steroid",
        "immunosuppression",
        "chemo",
        "antibiotic",
        "biologic",
    ),
    "comorbidity": (
        "comorbidity",
        "comorbid",
        "diabetes",
        "hypertension",
        "asthma",
        "copd",
        "autoimmune",
    ),
    "outcome": (
        "response",
        "responder",
        "nonresponder",
        "outcome",
        "survival",
        "pfs",
        "os",
        "recurrence",
        "mortality",
    ),
    "pregnancy": (
        "pregnancy",
        "pregnant",
        "trimester",
        "gestational",
        "gestation",
        "parity",
        "menopause",
    ),
    "assay_modality": (
        "assay",
        "modality",
        "atac",
        "adt",
        "cite",
        "vdj",
        "tcr",
        "bcr",
        "multiome",
        "spatial",
    ),
    "technology": (
        "platform",
        "technology",
        "chemistry",
        "kit",
        "chromium",
        "smartseq",
        "smart_seq",
        "dropseq",
        "drop_seq",
        "indrop",
    ),
    "batch": (
        "batch",
        "library",
        "lane",
        "flowcell",
        "run",
        "pool",
        "hashtag",
        "hto",
        "multiplex",
        "demux",
    ),
    "qc": (
        "ncount",
        "n_count",
        "nfeature",
        "n_feature",
        "percent_mt",
        "pct_mt",
        "mito",
        "doublet",
        "scrublet",
        "umis",
    ),
}

_SIDECAR_DIR = Path("yuj_setup/accessions")
# Override lets multiple yuj jobs share one code/venv checkout (remote_dir)
# without sharing a results/ folder -- the worker skips any item whose result
# file already exists, so a shared folder would make a rescan job silently
# skip everything a prior job (or an unrelated historical run) already wrote.
_OUT_DIR = Path(os.environ.get("SC_OUTPUT_DIR", "results"))

_DEFAULT_SOCKS = "10.195.102.16:1080"

# Opt into HTTPS for NCBI files when proxies intercept FTP. NCBI serves the
# same paths over HTTPS; default to FTP for compatibility.
_FORCE_HTTP = os.environ.get("SC_FORCE_HTTP", "") == "1"

# Cap accession runtime because downloads lack stall timeouts. Timed-out
# items record 0 and count as done.
_PER_ITEM_DEADLINE = int(os.environ.get("SC_ITEM_DEADLINE", "1800"))


class _StallTimeoutError(BaseException):
    """BaseException bypasses broad except Exception paths."""


def _on_alarm(signum: int, frame: object) -> None:  # noqa: ARG001
    raise _StallTimeoutError


def _read_proxy() -> str | None:
    pf = Path("socks_proxy")
    if not pf.exists():
        return _DEFAULT_SOCKS
    val = pf.read_text().strip()
    return None if not val or val.lower() == "direct" else val


def main() -> None:
    accession = sys.argv[1]
    _OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = _OUT_DIR / f"{accession}.json"
    if out_path.exists():
        return

    meta = json.loads((_SIDECAR_DIR / f"{accession}.json").read_text())

    signal.signal(signal.SIGALRM, _on_alarm)
    signal.alarm(_PER_ITEM_DEADLINE)
    try:
        result = process_accession(
            accession=meta["accession"],
            source=meta.get("source", ""),
            aliases=meta.get("aliases", ""),
            links=meta.get("links", []),
            output_dir=_OUT_DIR,
            chunk_size=_CHUNK_SIZE,
            retries=_RETRIES,
            timeout=_TIMEOUT,
            max_timeout=_MAX_TIMEOUT,
        )
    except _StallTimeoutError:
        logger.warning("stall-timeout (%ss) on %s", _PER_ITEM_DEADLINE, accession)
        result = {
            "accession": accession,
            "source": meta.get("source", ""),
            "aliases": meta.get("aliases", ""),
            "single_cell_count": 0,
            "single_cell_file_format": "",
            "stall_timeout": True,
        }
    except Exception:
        logger.exception("failed accession %s", accession)
        result = {"accession": accession, "error": True}
    finally:
        signal.alarm(0)

    # Write via temp file so resume sees only complete outputs.
    tmp = out_path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(result, indent=2))
    tmp.replace(out_path)


def process_accession(
    accession: str,
    source: str,
    aliases: str,
    links: list[tuple[str, str]],
    output_dir: Path,
    chunk_size: int,
    retries: int,
    timeout: int,
    max_timeout: int,
) -> dict[str, Any]:
    dl_links = [DownloadLink(type=FileType(t), url=u) for t, u in links]
    dl_links.sort(key=lambda x: FILE_TYPE_PRIORITY.index(x.type))

    acc_dir = output_dir / accession
    acc_dir.mkdir(parents=True, exist_ok=True)

    # rds/h5ad have no range-read path for metadata (pyreadr/rpy2/gzipped-h5ad
    # all need a local file), so count_link's full download below is the only
    # chance to read obs/meta.data columns before the file is deleted --
    # _scan_metadata used to re-download the same URL from scratch just to
    # read the same columns. Cache by URL here; keyed on URL (not accession)
    # since a study can have several rds/h5ad links.
    embedded_cols_cache: dict[str, list[str]] = {}

    full_download = functools.partial(
        download_file_ftp,
        chunk_size=chunk_size,
        num_retries=retries,
        timeout=timeout,
        max_wait=max_timeout,
        with_pbar=False,
        socks_proxy=_read_proxy(),
        max_bytes=_MAX_FILE_BYTES,
    )

    def count_link(  # noqa: PLR0911
        link: DownloadLink, *, force_http: bool = False
    ) -> tuple[int, int, str, str]:
        # EBI serves supplementary files over HTTPS only; ftp:// there times out
        use_http = force_http or _FORCE_HTTP or "ebi.ac.uk" in link.url
        ftp_url = link.url.replace("https://", "ftp://", 1)
        # ftp.ncbi.nlm.nih.gov serves the same paths over HTTPS; HTTP branches receive the converted URL.
        http_url = link.url.replace("ftp://", "https://", 1)

        if link.type in {FileType.H5, FileType.H5ad}:
            dims = count_h5_dims_via_http(link.url, link.type)
            if dims is not None:
                n, g = dims
                return (n, g, link.type.value, "") if n > 0 else (0, -1, "", "")

        if link.type == FileType.Tar and not use_http:
            n = count_tar_via_stream(
                link.url, chunk_size, timeout=timeout, socks_proxy=_read_proxy()
            )
            if n is not None:
                return (n, -1, link.type.value, "") if n > 0 else (0, -1, "", "")

        dest_path = acc_dir / link.extract_filename()
        try:
            if link.type in HEADER_STREAMABLE and not is_cell_major(link.url):
                head = _MTX_HEAD_BYTES if link.type == FileType.Mtx else chunk_size
                if use_http:
                    dest_path = download_http_head(
                        http_url,
                        dest_path,
                        max_decompressed=head,
                        num_retries=retries,
                        timeout=timeout,
                        max_wait=max_timeout,
                    )
                else:
                    dest_path = download_ftp_head(
                        ftp_url,
                        dest_path,
                        max_decompressed=head,
                        num_retries=retries,
                        timeout=timeout,
                        max_wait=max_timeout,
                        socks_proxy=_read_proxy(),
                    )
            else:
                if not _check_disk_space(output_dir):
                    logger.warning("skipping %s: insufficient disk space", link.url)
                    return (0, -1, "", "")
                if use_http:
                    dest_path = download_file(
                        http_url,
                        dest_path,
                        chunk_size=chunk_size,
                        num_workers=1,
                        num_retries=retries,
                        timeout=timeout,
                        max_wait=max_timeout,
                        with_pbar=False,
                        validate_content_length=False,
                    )
                else:
                    dest_path = full_download(url=ftp_url, dest_path=dest_path)
        except FileTooLargeError as exc:
            logger.info("skipping oversized file: %s", exc)
            return (0, -1, "", "")
        except Exception:
            logger.exception("failed to download %s", link.url)
            return (0, -1, "", "")

        try:
            result = _resolve_file_type(dest_path, link.type, chunk_size)
            if result is None:
                logger.debug("unable to find file type for %s", dest_path)
                return (0, -1, "", "")
            file_type, sniffed_path = result
            if file_type == FileType.Tar:
                # Count genes per member; the enclosing tar has no gene axis.
                n, g, fmt = process_tar_with_genes(sniffed_path, chunk_size)
            else:
                n = _get_sc_count(sniffed_path, file_type, chunk_size)
                g = gene_count_for(sniffed_path, file_type, chunk_size) if n > 0 else -1
                fmt = file_type.value
            if file_type == FileType.Rds:
                embedded_cols_cache[link.url] = rds_columns(sniffed_path)
            elif file_type == FileType.H5ad:
                embedded_cols_cache[link.url] = h5ad_obs_columns_local(sniffed_path)
            if sniffed_path != dest_path:
                sniffed_path.unlink(missing_ok=True)
            rds = last_rds_class() if file_type == FileType.Rds else ""
            return (n, g, fmt, rds) if n > 0 else (0, -1, "", "")  # noqa: TRY300
        except RdsReadError as exc:
            _record_rds_failure(output_dir, accession, link.url, exc.reason)
            return (0, -1, "", "")
        except Exception:
            logger.exception("failed to process %s", link.url)
            return (0, -1, "", "")
        finally:
            dest_path.unlink(missing_ok=True)
            dest_path.with_suffix("").unlink(missing_ok=True)

    samples: dict[str, dict[str, Any]] = {}

    def _merge(
        key: str,
        n_cells: int,
        n_genes: int,
        fmt: str,
        rds_class: str,
        nnz: int,
        src: str,
    ) -> None:
        if n_cells <= 0:
            return
        s = samples.setdefault(
            key,
            {
                "sample": key,
                "cells": 0,
                "genes": -1,
                "nnz": -1,
                "cells_format": "",
                "genes_format": "",
                "rds_class": "",
                "unfiltered": False,
            },
        )
        if n_cells > s["cells"]:
            s["cells"] = n_cells
            s["cells_format"] = fmt
            s["unfiltered"] = _is_unfiltered(n_cells, nnz, n_genes, src)
        if n_genes > s["genes"]:
            s["genes"] = n_genes
            s["genes_format"] = fmt
        if nnz > 0:
            s["nnz"] = nnz
        if rds_class:
            s["rds_class"] = rds_class

    for link in dl_links:
        # Stream per-GSM matrices separately to preserve sample attribution.
        # SC_FORCE_HTTP skips this FTP path; HTTPS handling below does the same work.
        if (
            link.type == FileType.Tar
            and "ebi.ac.uk" not in link.url
            and not _FORCE_HTTP
        ):
            members = count_tar_members_via_stream(
                link.url, chunk_size, timeout=timeout, socks_proxy=_read_proxy()
            )
            if members is not None:
                for mname, m_cells, m_genes, m_fmt, m_nnz in members:
                    # GSM may be in the member name (multi-sample tar) or only in the
                    # tar filename (per-GSM 10x triplet tar)
                    key = _sample_key(mname, accession)
                    if key == accession:
                        key = _sample_key(link.url, accession)
                    _merge(key, m_cells, m_genes, m_fmt, "", m_nnz, mname)
                continue
        # Retry FTP zeroes and exceptions over HTTPS: blocked transfers can look
        # like empty files or raise before the cell-count check.
        retryable = not _FORCE_HTTP and "ebi.ac.uk" not in link.url
        try:
            n_cells, n_genes, fmt, rds_class = count_link(link)
        except Exception:
            if not retryable:
                raise
            logger.debug("ftp path raised for %s; retrying over https", link.url)
            n_cells, n_genes, fmt, rds_class = 0, -1, "", ""
        if n_cells <= 0 and retryable:
            try:
                n_cells, n_genes, fmt, rds_class = count_link(link, force_http=True)
            except Exception:
                logger.debug("https retry also failed for %s", link.url)
                n_cells, n_genes, fmt, rds_class = 0, -1, "", ""
        nnz = last_mtx_nnz() if fmt == FileType.Mtx.value else -1
        _merge(
            _sample_key(link.url, accession),
            n_cells,
            n_genes,
            fmt,
            rds_class,
            nnz,
            link.url,
        )

    # Keep counts if the optional metadata scan exceeds the deadline.
    try:
        metadata = _scan_metadata(
            links,
            accession,
            acc_dir,
            retries=retries,
            timeout=timeout,
            max_timeout=max_timeout,
            embedded_cols_cache=embedded_cols_cache,
        )
    except _StallTimeoutError:
        logger.warning("metadata scan hit deadline on %s; counts kept", accession)
        metadata = {}

    shutil.rmtree(acc_dir, ignore_errors=True)

    return _build_result(accession, source, aliases, samples, metadata)


def _is_unfiltered(cells: int, nnz: int, genes: int, src: str) -> bool:
    if cells in _WHITELIST_SIZES:
        return True
    if 0 < genes < _MIN_NNZ_PER_CELL:
        # HTO/ADT libraries have few features; use relative density to distinguish
        # filtered matrices from mostly empty raw droplets.
        return nnz > 0 and cells > 0 and nnz < cells * genes * _MIN_LOW_FEATURE_DENSITY
    if nnz > 0 and cells > 0 and nnz < cells * _MIN_NNZ_PER_CELL:
        return True
    low = src.rsplit("/", 1)[-1].lower()
    return "raw" in low and "filtered" not in low and cells > 500_000  # noqa: PLR2004


def _sample_key(url: str, accession: str) -> str:
    # Match basenames: GEO directory stubs like GSMxxxxnnn can produce colliding IDs.
    m = re.search(r"GSM\d+", url.rsplit("/", 1)[-1])
    if m:
        return m.group(0)
    for seg in url.split("/"):
        if re.fullmatch(r"GSM\d+", seg):
            return seg
    return accession


def _is_metadata_link(url: str) -> bool:
    # A supplementary csv/tsv/txt/xlsx that isn't an obvious raw-matrix
    # companion (matrix/barcodes/features/h5) is a metadata-table candidate
    # regardless of its filename -- real GEO deposits name their per-cell
    # annotation table all sorts of things ("Table_S3.xlsx", "cell_info.csv",
    # "SampleSheet.tsv") that don't contain any of _META_NAME_HINTS. Filtering
    # on filename first meant most real metadata tables were never even
    # opened; the actual gate is `_header_info`'s `per_cell` check (requires a
    # barcode-like index/column AND a recognized field), which is what stops
    # this from probing every DEG table or README as if it might be per-cell
    # metadata.
    name = url.rsplit("/", 1)[-1].lower().removesuffix(".gz")
    if any(x in name for x in _META_EXCLUDE):
        return False
    return name.endswith(_META_EXTS)


def _parse_header(line: str, name: str) -> list[str]:
    line = line.rstrip("\r\n")
    low = name.lower().removesuffix(".gz")
    if low.endswith(".csv"):
        delim = ","
    elif low.endswith(".tsv"):
        delim = "\t"
    else:
        # Infer the delimiter; R write.table often uses spaces in .txt files.
        delim = max(("\t", ",", ";", " "), key=line.count)
    return [c.strip().strip('"').strip("'") for c in line.split(delim)]


@functools.cache
def _squash(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", s)


def _match_field(col: str, hints: tuple[str, ...]) -> bool:
    # Token matching lets age match age_years, not percentage. The substring
    # fallback squashes separators on both sides so a multi-word hint matches
    # regardless of how the column joins its words -- "cell_type" (a
    # programmatic scanpy/Seurat obs name) and "Cell type" (a human-written
    # CSV header) squash to the same "celltype" and must match identically;
    # comparing `h in low` without squashing only ever matched the former.
    low = col.lower()
    toks = set(re.split(r"[^a-z0-9]+", low))
    squashed = _squash(low)
    for h in hints:
        if h in toks:
            return True
        h_squashed = _squash(h)
        if len(h_squashed) >= 8 and h_squashed in squashed:  # noqa: PLR2004
            return True
    return False


def _detect_field_map(cols: list[str]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for cat, hints in _META_FIELD_HINTS.items():
        matched = [c for c in cols if _match_field(c, hints)]
        if matched:
            out[cat] = matched
    return out


def _cell_index_type(value: str) -> str:
    v = value.strip().strip('"').strip("'")
    if re.search(r"[ACGTNacgtn]{10,}", v):
        return "barcode"
    if re.fullmatch(r"(?i)cell[\s_.\-]?\d+", v):
        return "renamed"
    return "id"


def _header_info_from_rows(cols: list[str], row: list[str], name: str) -> dict[str, Any]:
    low = [c.lower() for c in cols]
    named_barcode = any(h in c for c in low for h in _BARCODE_COL_HINTS)
    # R row indexes use an empty header or leave data rows one column wider.
    has_unnamed_index = bool(row) and (
        (bool(cols) and cols[0] == "") or len(row) > len(cols)
    )
    cell_index = _cell_index_type(row[0]) if has_unnamed_index else None
    fmap = _detect_field_map(cols)
    fields = sorted(fmap)
    has_barcode = named_barcode or cell_index == "barcode"
    return {
        "file": name,
        "columns": cols[:_MAX_META_COLS],
        "n_columns": len(cols),
        "has_barcode": has_barcode,
        "cell_index": cell_index,
        "has_celltype": "celltype" in fmap,
        "has_donor": "subject" in fmap,
        "has_demographics": any(c in fmap for c in _DEMOGRAPHIC_CATEGORIES),
        "fields": fields,
        "category_to_columns": {k: v[:6] for k, v in fmap.items()},
        "per_cell": (has_barcode or cell_index is not None) and bool(fields),
    }


def _header_info(first: str, second: str, name: str) -> dict[str, Any]:
    cols = _parse_header(first, name)
    row = _parse_header(second, name) if second.strip() else []
    return _header_info_from_rows(cols, row, name)


def _columns_info(cols: list[str], name: str, fmt: str) -> dict[str, Any]:
    """Same shape as _header_info, for formats whose columns come from a real
    parser (h5ad obs, rds meta.data/colData) instead of a text header row --
    every row in those tables is one cell by construction, so there is no
    barcode/index heuristic to run.
    """
    fmap = _detect_field_map(cols)
    fields = sorted(fmap)
    return {
        "file": name,
        "columns": cols[:_MAX_META_COLS],
        "n_columns": len(cols),
        "has_barcode": None,
        "cell_index": None,
        "has_celltype": "celltype" in fmap,
        "has_donor": "subject" in fmap,
        "has_demographics": any(c in fmap for c in _DEMOGRAPHIC_CATEGORIES),
        "fields": fields,
        "category_to_columns": {k: v[:6] for k, v in fmap.items()},
        "per_cell": bool(cols),
        "embedded": True,
        "format": fmt,
    }


def _download_for_probe(
    url: str,
    dest: Path,
    *,
    retries: int,
    timeout: int,
    max_timeout: int,
    proxy: str | None,
    max_bytes: int,
) -> Path:
    """Shared HTTP-vs-FTP download branch for the h5ad/rds/xlsx metadata probes."""
    use_http = "ebi.ac.uk" in url or _FORCE_HTTP
    if use_http:
        return download_file(
            url.replace("ftp://", "https://", 1),
            dest,
            chunk_size=_CHUNK_SIZE,
            num_workers=1,
            num_retries=retries,
            timeout=timeout,
            max_wait=max_timeout,
            with_pbar=False,
            validate_content_length=False,
        )
    return download_file_ftp(
        url.replace("https://", "ftp://", 1),
        dest,
        chunk_size=_CHUNK_SIZE,
        num_retries=retries,
        timeout=timeout,
        max_wait=max_timeout,
        socks_proxy=proxy,
        max_bytes=max_bytes,
    )


def _probe_h5ad_metadata(
    url: str,
    acc_dir: Path,
    *,
    retries: int,
    timeout: int,
    max_timeout: int,
    cached_cols: list[str] | None = None,
) -> dict[str, Any] | None:
    """Real obs-column check for an .h5ad.

    Tries HTTP range reads first (cheap: HDF5 group listing needs only
    metadata, not the file's actual data) -- but most GEO .h5ad deposits are
    .h5ad.gz, and gzip-wrapped HDF5 has no random access, so that always
    fails for them. Falls back to a full download (same _resolve_file_type
    decompression the counting path already uses for exactly this reason).

    `cached_cols` short-circuits both: set when count_link already had to
    fully download this same URL for the cell count and read its obs columns
    while the file was local, so this probe reuses that instead of
    downloading the same file a second time. Empty list means it tried and
    found no columns -- also a reason to skip re-downloading, not to retry.
    """
    if cached_cols is not None:
        if not cached_cols:
            return None
        info = _columns_info(cached_cols, url.rsplit("/", 1)[-1], "h5ad")
        info["url"] = url
        return info

    cols = h5ad_obs_columns_via_http(url)
    if cols is not None:
        info = _columns_info(cols, url.rsplit("/", 1)[-1], "h5ad")
        info["url"] = url
        return info

    name = url.rsplit("/", 1)[-1]
    dest = acc_dir / f"h5ad_meta_{name}"
    try:
        path = _download_for_probe(
            url,
            dest,
            retries=retries,
            timeout=timeout,
            max_timeout=max_timeout,
            proxy=_read_proxy(),
            max_bytes=_MAX_FILE_BYTES,
        )
        resolved = _resolve_file_type(path, FileType.H5ad, _CHUNK_SIZE)
        if resolved is None:
            return None
        _, sniffed = resolved
        cols = h5ad_obs_columns_local(sniffed)
    except Exception:
        logger.debug("h5ad metadata probe failed for %s", url)
        return None
    finally:
        dest.unlink(missing_ok=True)
        dest.with_suffix("").unlink(missing_ok=True)
    if not cols:
        return None
    info = _columns_info(cols, name, "h5ad")
    info["url"] = url
    return info


def _probe_rds_metadata(
    url: str,
    acc_dir: Path,
    *,
    retries: int,
    timeout: int,
    max_timeout: int,
    cached_cols: list[str] | None = None,
) -> dict[str, Any] | None:
    """Real column check for an .rds's data.frame / meta.data / colData.

    pyreadr/rpy2 need a local file (no range-read support for R's serialization
    format). `cached_cols` reuses the columns count_link already read from
    its own full download of this URL for the cell count, instead of
    downloading the same file again here. Empty list means that read found
    no columns -- also a reason to skip, not to retry.
    """
    if cached_cols is not None:
        if not cached_cols:
            return None
        name = url.rsplit("/", 1)[-1]
        info = _columns_info(cached_cols, name, "rds")
        info["url"] = url
        rds_class = last_rds_class()
        if rds_class:
            info["rds_class"] = rds_class
        return info

    name = url.rsplit("/", 1)[-1]
    dest = acc_dir / f"rds_meta_{name}"
    try:
        path = _download_for_probe(
            url,
            dest,
            retries=retries,
            timeout=timeout,
            max_timeout=max_timeout,
            proxy=_read_proxy(),
            max_bytes=_MAX_FILE_BYTES,
        )
        resolved = _resolve_file_type(path, FileType.Rds, _CHUNK_SIZE)
        if resolved is None:
            return None
        _, sniffed = resolved
        cols = rds_columns(sniffed)
    except Exception:
        logger.debug("rds metadata probe failed for %s", url)
        return None
    finally:
        dest.unlink(missing_ok=True)
        dest.with_suffix("").unlink(missing_ok=True)
    if not cols:
        return None
    info = _columns_info(cols, name, "rds")
    info["url"] = url
    rds_class = last_rds_class()
    if rds_class:
        info["rds_class"] = rds_class
    return info


def _probe_metadata(
    url: str,
    acc_dir: Path,
    *,
    retries: int,
    timeout: int,
    max_timeout: int,
    proxy: str | None,
) -> dict[str, Any] | None:
    name = url.rsplit("/", 1)[-1]
    dest = acc_dir / f"meta_{name}"
    use_http = "ebi.ac.uk" in url
    try:
        if use_http:
            path = download_http_head(
                url,
                dest,
                max_decompressed=_META_HEAD_BYTES,
                num_retries=retries,
                timeout=timeout,
                max_wait=max_timeout,
            )
        else:
            path = download_ftp_head(
                url.replace("https://", "ftp://", 1),
                dest,
                max_decompressed=_META_HEAD_BYTES,
                num_retries=retries,
                timeout=timeout,
                max_wait=max_timeout,
                socks_proxy=proxy,
            )
        with path.open("rb") as f:
            first = f.readline(1 << 20).decode("utf-8", "replace")
            second = f.readline(1 << 20).decode("utf-8", "replace")
    except Exception:
        logger.debug("metadata probe failed for %s", url)
        return None
    finally:
        dest.unlink(missing_ok=True)
        dest.with_suffix("").unlink(missing_ok=True)

    info = _header_info(first, second, name)
    info["url"] = url
    return info


_MAX_XLSX_BYTES = 50 * 1024 * 1024


def _xlsx_header_rows(path: Path) -> tuple[list[str], list[str]] | None:
    """First two rows of an .xlsx's first sheet, as strings. None on any
    failure -- including a legacy .xls, which openpyxl cannot open (no xlrd
    dependency here; .xls is rare enough in current GEO deposits to skip
    rather than add a second Excel-reading library for it).
    """
    try:
        import openpyxl
    except ImportError:
        return None
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(max_row=2, values_only=True)
            first = next(it, None)
            second = next(it, None)
        finally:
            wb.close()
    except Exception:
        return None
    if first is None:
        return None
    cols = ["" if c is None else str(c).strip() for c in first]
    row = ["" if c is None else str(c).strip() for c in second] if second else []
    return cols, row


def _probe_xlsx_metadata(
    url: str,
    acc_dir: Path,
    *,
    retries: int,
    timeout: int,
    max_timeout: int,
    proxy: str | None,
) -> dict[str, Any] | None:
    # xlsx is a zip archive with its central directory at the end, so a
    # byte-range head (as _probe_metadata does for csv/tsv/txt) would produce
    # an unparseable truncated file. Metadata tables are small, so download in
    # full instead, capped well under a real expression matrix's size.
    name = url.rsplit("/", 1)[-1]
    dest = acc_dir / f"xlsx_meta_{name}"
    try:
        path = _download_for_probe(
            url,
            dest,
            retries=retries,
            timeout=timeout,
            max_timeout=max_timeout,
            proxy=proxy,
            max_bytes=_MAX_XLSX_BYTES,
        )
        if path.stat().st_size > _MAX_XLSX_BYTES:
            return None
        rows = _xlsx_header_rows(path)
    except Exception:
        logger.debug("xlsx metadata probe failed for %s", url)
        return None
    finally:
        dest.unlink(missing_ok=True)
    if rows is None:
        return None
    cols, row = rows
    info = _header_info_from_rows(cols, row, name)
    info["url"] = url
    return info


def _scan_tar_metadata(
    tar_url: str,
    accession: str,
    *,
    timeout: int,
    proxy: str | None,
) -> dict[str, dict[str, Any]]:
    # Keep tar URL and member path for retrieval via tar -xOf; cap members
    # because streamed tars lack random access.
    ftp_url = tar_url.replace("https://", "ftp://", 1)
    by_key: dict[str, dict[str, Any]] = {}
    probed = 0
    seen = 0
    try:
        with (
            ftp_data_stream(ftp_url, timeout=timeout, socks_proxy=proxy) as stream,
            tarfile.open(fileobj=stream, mode="r|*") as tar,
        ):
            for member in tar:
                if probed >= _MAX_TAR_META_MEMBERS or seen >= _MAX_TAR_MEMBERS_WALKED:
                    break
                seen += 1
                if not member.isfile() or not _is_metadata_link(member.name):
                    continue
                f = tar.extractfile(member)
                if f is None:
                    continue
                head = _member_head(f, member.name, _META_HEAD_BYTES)
                text = head.decode("utf-8", "replace")
                nl = text.find("\n")
                first, second = (text, "") if nl == -1 else (text[:nl], text[nl + 1 :])
                base = member.name.rsplit("/", 1)[-1]
                info = _header_info(first, second, base)
                if not info["per_cell"]:
                    probed += 1
                    continue
                info["url"] = tar_url
                info["tar_member"] = member.name
                key = _sample_key(member.name, accession)
                prev = by_key.get(key)
                if prev is None or len(info["fields"]) > len(prev["fields"]):
                    by_key[key] = info
                probed += 1
    except Exception:
        logger.debug("tar metadata scan failed for %s", tar_url)
    return by_key


def _scan_metadata(
    links: list[tuple[str, str]],
    accession: str,
    acc_dir: Path,
    *,
    retries: int,
    timeout: int,
    max_timeout: int,
    embedded_cols_cache: dict[str, list[str]] | None = None,
) -> dict[str, dict[str, Any]]:
    proxy = _read_proxy()
    by_key: dict[str, dict[str, Any]] = {}
    embedded_cols_cache = embedded_cols_cache or {}

    def _keep(key: str, info: dict[str, Any]) -> None:
        if not info["per_cell"]:
            return
        prev = by_key.get(key)
        if prev is None or len(info["fields"]) > len(prev["fields"]):
            by_key[key] = info

    tar_urls: list[str] = []
    for t, url in links:
        if t == FileType.Tar.value:
            tar_urls.append(url)
        if t == FileType.H5ad.value:
            info = _probe_h5ad_metadata(
                url,
                acc_dir,
                retries=retries,
                timeout=timeout,
                max_timeout=max_timeout,
                cached_cols=embedded_cols_cache.get(url),
            )
            if info is not None:
                _keep(_sample_key(url, accession), info)
            continue
        if t == FileType.Rds.value:
            info = _probe_rds_metadata(
                url,
                acc_dir,
                retries=retries,
                timeout=timeout,
                max_timeout=max_timeout,
                cached_cols=embedded_cols_cache.get(url),
            )
            if info is not None:
                _keep(_sample_key(url, accession), info)
            continue
        if not _is_metadata_link(url):
            continue
        name = url.rsplit("/", 1)[-1].lower().removesuffix(".gz")
        if name.endswith(_META_TABLE_EXTS):
            info = _probe_xlsx_metadata(
                url,
                acc_dir,
                retries=retries,
                timeout=timeout,
                max_timeout=max_timeout,
                proxy=proxy,
            )
        else:
            info = _probe_metadata(
                url,
                acc_dir,
                retries=retries,
                timeout=timeout,
                max_timeout=max_timeout,
                proxy=proxy,
            )
        if info is not None:
            _keep(_sample_key(url, accession), info)

    if not by_key:
        for tar_url in tar_urls[:_MAX_TARS_SCANNED]:
            found = _scan_tar_metadata(tar_url, accession, timeout=timeout, proxy=proxy)
            for key, info in found.items():
                _keep(key, info)

    return by_key


def _build_result(
    accession: str,
    source: str,
    aliases: str,
    samples: dict[str, dict[str, Any]],
    metadata: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    # single_cell_count = max(per-GSM sum, largest study-level file), avoiding
    # double counts when combined and per-sample matrices coexist.
    gsm = [s for s in samples.values() if s["sample"].startswith("GSM")]
    study = [s for s in samples.values() if not s["sample"].startswith("GSM")]
    gsm_sum = sum(s["cells"] for s in gsm)
    study_max = max((s["cells"] for s in study), default=0)

    rows = sorted(gsm, key=lambda s: s["cells"], reverse=True)
    if not gsm or study_max > gsm_sum:
        rows += sorted(study, key=lambda s: s["cells"], reverse=True)

    proj_meta = metadata.get(accession)

    def _fmt(s: dict[str, Any]) -> str:
        return s["genes_format"] if s["genes"] > 0 else s["cells_format"]

    def _sample_meta(s: dict[str, Any]) -> dict[str, Any] | None:
        ext = metadata.get(s["sample"])
        if ext is not None:
            return ext
        # _scan_metadata already probes every H5ad/Rds link's real columns
        # (see _probe_h5ad_metadata / _probe_rds_metadata) and would have
        # landed here via the lookup above if that succeeded. Reaching this
        # branch means the probe failed or wasn't attempted for this sample
        # (e.g. a stall-timeout, or this sample's file never resolved to a
        # link) -- unknown, not a positive signal, so no has_celltype claim.
        if _fmt(s) in _EMBEDDED_METADATA_FORMATS:
            m: dict[str, Any] = {
                "embedded": True,
                "format": _fmt(s),
                "has_celltype": False,
                "unverified": True,
            }
            if s["rds_class"]:
                m["rds_class"] = s["rds_class"]
            return m
        if s["sample"].startswith("GSM") and proj_meta is not None:
            return {**proj_meta, "scope": "study"}
        return None

    def _scope(s: dict[str, Any]) -> str:
        return "sample" if s["sample"].startswith("GSM") else "study"

    per_sample = [
        {
            "sample": s["sample"],
            "scope": _scope(s),
            "cells": s["cells"],
            "genes": s["genes"] if s["genes"] > 0 else None,
            "nnz": s["nnz"] if s["nnz"] > 0 else None,
            "file_format": _fmt(s),
            "unfiltered": s["unfiltered"],
            "metadata": _sample_meta(s),
        }
        for s in rows
    ]

    all_meta = [p["metadata"] for p in per_sample if p["metadata"]]
    if proj_meta:
        all_meta.append(proj_meta)
    meta_fields = sorted({f for m in all_meta for f in m.get("fields", [])})

    out: dict[str, Any] = {
        "accession": accession,
        "source": source,
        "aliases": aliases,
        "single_cell_count": max(gsm_sum, study_max),
        "single_cell_file_format": _fmt(rows[0]) if rows else "",
        "n_samples": len(per_sample),
        "n_gsm_samples": sum(1 for p in per_sample if p["scope"] == "sample"),
        "unfiltered": any(p["unfiltered"] for p in per_sample),
        "has_metadata": bool(metadata) or any(p["metadata"] for p in per_sample),
        "has_celltype": any(m.get("has_celltype") for m in all_meta),
        "has_donor": any(m.get("has_donor") for m in all_meta),
        "has_demographics": any(m.get("has_demographics") for m in all_meta),
        "metadata_fields": meta_fields,
        "project_metadata": proj_meta,
        "per_sample": per_sample,
    }
    if rows and rows[0]["rds_class"]:
        out["rds_class"] = rows[0]["rds_class"]
    return out


if __name__ == "__main__":
    main()
