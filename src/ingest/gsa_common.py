"""Shared HTTP client, xlsx parsing, downloads, and SQLite writers for GSA stages."""

from __future__ import annotations

import datetime
import io
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import TYPE_CHECKING, Any

import httpx
import openpyxl
from tqdm import tqdm

if TYPE_CHECKING:
    import sqlite3
    from collections.abc import Callable
    from pathlib import Path

logger = logging.getLogger(__name__)

GSA_BASE_URL = "https://ngdc.cncb.ac.cn"
GSA_DOWNLOAD_URL = "https://download.cncb.ac.cn"

# GSA sits behind a WAF that drops connections without a browser User-Agent.
BROWSER_UA = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)


def make_client(timeout: int) -> httpx.Client:
    return httpx.Client(
        headers={"User-Agent": BROWSER_UA},
        timeout=timeout,
        follow_redirects=True,
    )


def parse_workbook(xlsx_bytes: bytes) -> dict[str, list[dict]]:
    """Parse an xlsx into {sheet_title: [row_dict, ...]}."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=False)
    try:
        sheets: dict[str, list[dict]] = {}
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header = [
                    " ".join(str(h).split()) if h is not None else ""
                    for h in next(rows)
                ]
            except StopIteration:
                sheets[ws.title] = []
                continue
            records = [
                {header[i]: row[i] for i in range(min(len(header), len(row)))}
                for row in rows
                if any(v is not None for v in row)
            ]
            sheets[ws.title] = records
        return sheets
    finally:
        wb.close()


def download_studies(
    accessions: list[str],
    out_dir: Path,
    fetch: Callable[[str], bytes | None],
    num_workers: int,
    desc: str,
    suffix: str = ".xlsx",
) -> dict[str, str]:
    """Fetch each study's file (workbook/manifest) to disk; acc -> status."""
    out_dir.mkdir(parents=True, exist_ok=True)
    statuses: dict[str, str] = {}

    def _worker(acc: str) -> tuple[str, str]:
        dest = out_dir / f"{acc}{suffix}"
        if dest.exists() and dest.stat().st_size > 0:
            return acc, "cached"
        try:
            data = fetch(acc)
        except Exception:
            logger.exception("failed to fetch %s", acc)
            return acc, "failed"
        if not data:
            return acc, "empty"
        tmp = dest.with_name(dest.name + ".part")
        tmp.write_bytes(data)
        tmp.rename(dest)  # atomic: a partial download never looks complete
        return acc, "ok"

    with (
        ThreadPoolExecutor(num_workers) as pool,
        tqdm(total=len(accessions), unit="study", desc=desc) as pbar,
    ):
        futures = [pool.submit(_worker, a) for a in accessions]
        for f in as_completed(futures):
            acc, status = f.result()
            statuses[acc] = status
            pbar.update(1)
    return statuses


def coerce(value: Any) -> Any:
    """Normalise a cell value for SQLite: blank -> None, non-primitive -> str."""
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


def to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def map_row(
    record: dict,
    header_to_column: dict[str, str],
    columns: list[str],
    static: dict[str, Any] | None = None,
) -> tuple:
    """Build a column-ordered tuple from a sheet row, keyed by header name."""
    out: dict[str, Any] = dict(static or {})
    for header, column in header_to_column.items():
        out[column] = coerce(record.get(header))
    return tuple(out.get(column) for column in columns)


def date_filter_active(start_date: datetime.date, end_date: datetime.date) -> bool:
    """Whether a real date filter was requested; today..today means full sync."""
    today = datetime.datetime.now(datetime.UTC).date()
    return not (start_date == end_date == today)


def in_range(
    release_time: str | None,
    start_date: datetime.date,
    end_date: datetime.date,
) -> bool:
    if not release_time:
        return False
    try:
        released = datetime.date.fromisoformat(release_time[:10])
    except ValueError:
        return False
    return start_date <= released <= end_date


def apply_write_pragmas(conn: sqlite3.Connection) -> None:
    """Bulk-load PRAGMAs."""
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=OFF")
    conn.execute("PRAGMA cache_size=-800000")


def bulk_upsert(
    conn: sqlite3.Connection,
    table: str,
    columns: list[str],
    rows: list[tuple],
) -> int:
    if not rows:
        return 0
    cols = ", ".join(columns)
    placeholders = ", ".join("?" * len(columns))
    sql = f"INSERT OR REPLACE INTO {table} ({cols}) VALUES ({placeholders})"  # noqa: S608
    conn.executemany(sql, rows)
    return len(rows)


def bulk_upsert_preserve(
    conn: sqlite3.Connection,
    table: str,
    columns: list[str],
    rows: list[tuple],
    conflict_col: str,
) -> int:
    """Upsert writing only columns, preserving columns owned by other stages."""
    if not rows:
        return 0
    cols = ", ".join(columns)
    placeholders = ", ".join("?" * len(columns))
    updates = ", ".join(f"{c}=excluded.{c}" for c in columns if c != conflict_col)
    sql = (
        f"INSERT INTO {table} ({cols}) VALUES ({placeholders}) "  # noqa: S608
        f"ON CONFLICT({conflict_col}) DO UPDATE SET {updates}"
    )
    conn.executemany(sql, rows)
    return len(rows)
